import os
import pytest
import pymupdf as fitz
import docx
from app.documents.pdf_parser import extract_text_from_pdf
from app.documents.docx_parser import extract_text_from_docx
from app.documents.downloader import save_temp_file, validate_file, cleanup_file
from app.documents.extractor import extract_text_from_file


def test_pdf_parser(tmp_path):
    pdf_file = tmp_path / "sample_resume.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "John Doe\nPython Developer\nExperience: 3 years")
    doc.save(str(pdf_file))
    doc.close()

    extracted = extract_text_from_pdf(str(pdf_file))
    assert "John Doe" in extracted
    assert "Python Developer" in extracted


def test_docx_parser(tmp_path):
    docx_file = tmp_path / "sample_resume.docx"
    doc = docx.Document()
    doc.add_paragraph("Jane Smith")
    doc.add_paragraph("Senior Backend Engineer")
    doc.save(str(docx_file))

    extracted = extract_text_from_docx(str(docx_file))
    assert "Jane Smith" in extracted
    assert "Senior Backend Engineer" in extracted


def test_downloader_and_validator():
    with pytest.raises(ValueError, match="Unsupported file format"):
        validate_file("invalid.exe", 1024)

    ext = validate_file("resume.pdf", 1024)
    assert ext == ".pdf"

    temp_path = save_temp_file(b"Sample PDF content", "sample.pdf")
    assert os.path.exists(temp_path)
    cleanup_file(temp_path)
    assert not os.path.exists(temp_path)
