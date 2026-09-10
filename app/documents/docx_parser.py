import docx
import logging

logger = logging.getLogger(__name__)


def extract_text_from_docx(file_path: str) -> str:
    """
    Extract raw text from a DOCX file using python-docx.
    """
    text_chunks = []
    try:
        doc = docx.Document(file_path)
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_chunks.append(paragraph.text.strip())
        
        # Also extract table text if any
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    text_chunks.append(row_text)
    except Exception as e:
        logger.error(f"Error extracting text from DOCX file {file_path}: {e}")
        raise ValueError(f"Failed to parse DOCX document: {str(e)}")

    full_text = "\n\n".join(text_chunks)
    if not full_text.strip():
        raise ValueError("DOCX document contains no readable text.")
    
    return full_text
